"""Outreach router — import targets from Excel, one-click cold email send."""
import os
import asyncio
import tempfile
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.database import get_db, DATABASE_PATH
from app.models import Contact, Firm, User, EntityType
from app.auth import get_current_user, require_admin
from app.email_utils import send_email, personalize_with_ai, is_smtp_configured

router = APIRouter(prefix="/outreach", tags=["outreach"])


# ── Hard-coded PI Law Firm cold email template ──

PI_EMAIL_SUBJECT = "{contact_name} — New Channel Partners"
PI_EMAIL_BODY = """Hey {contact_name},

I'm reaching out on behalf of New Channel Partners, a merchant bank based in Miami / New York. We're spending time in the personal injury law space and looking to connect with established firms like {firm}.

We've worked with PI law firms this past year, helping buy and sell 5 firms, and recently completed an $80M capital raise for the space.

Curious if you or the team have thought about strategic investors — either to help build the business with outside capital, or to plan an exit strategy for the near future.

Would you be open to a 30-minute intro call? I can connect you with our senior partners, Andrew and Will, who are available Tuesday or Thursday.

Thanks, and have a great week.
Matt with New Channel Partners"""


# ── Excel Import ──

@router.post("/import-excel")
async def import_excel(
    file: UploadFile = File(...),
    _admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    """Import PI Law Firm Outreach Tracker.xlsx — merges data from all sheets."""
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(500, "openpyxl not installed")

    if not file.filename or not file.filename.endswith('.xlsx'):
        raise HTTPException(400, "Please upload an .xlsx file")

    contents = await file.read()

    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
        tmp.write(contents)
        tmp_path = tmp.name

    try:
        wb = openpyxl.load_workbook(tmp_path, data_only=True)
    finally:
        os.unlink(tmp_path)

    created = 0
    skipped = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # Detect sheet format by headers
        headers = [str(ws.cell(1, c).value or "").strip().lower() for c in range(1, min(ws.max_column + 1, 30))]
        # Also check row 2 for headers (Master Outreach uses row 3)
        headers_r2 = [str(ws.cell(2, c).value or "").strip().lower() for c in range(1, min(ws.max_column + 1, 15))]
        headers_r3 = [str(ws.cell(3, c).value or "").strip().lower() for c in range(1, min(ws.max_column + 1, 15))]

        # Determine which row has headers and the mapping
        if "company name" in headers or "contact full name" in headers:
            # Format: Northeast / Westcoast style
            name_col = _find_col(headers, "contact full name", "contact primary phone number")
            firm_col = _find_col(headers, "company name")
            email_col = _find_col(headers, "contact primary e-mail")
            phone_col = _find_col(headers, "contact primary phone number")
            title_col = _find_col(headers, "contact title")
            region_col = _find_col(headers, "company region")
            start_row = 2

        elif "firm name" in headers_r3:
            # Format: Master Outreach
            name_col = 3  # Contact Name
            firm_col = 1  # Firm Name
            email_col = 5  # Contact Email
            phone_col = 4  # Contact Phone
            title_col = -1
            region_col = 2  # Region
            start_row = 4

        elif "firm" in headers or "contact" in headers:
            # Format: Southeast - PI Matt style
            firm_col = _find_col(headers, "firm")
            name_col = _find_col(headers, "contact")
            email_col = _find_col(headers, "email")
            phone_col = _find_col(headers, "phone")
            title_col = -1
            region_col = -1
            start_row = 2

        else:
            continue  # skip sheets we can't parse

        # Parse each row
        for r in range(start_row, ws.max_row + 1):
            name = str(ws.cell(r, (name_col if name_col > 0 else 1)).value or "").strip()
            if not name or len(name) < 2:
                continue

            email = str(ws.cell(r, (email_col if email_col > 0 else 1)).value or "").strip()
            if not email or "@" not in email:
                continue

            firm_name = str(ws.cell(r, (firm_col if firm_col > 0 else 1)).value or "").strip()
            phone = str(ws.cell(r, (phone_col if phone_col > 0 else 1)).value or "").strip()
            title = str(ws.cell(r, (title_col if title_col > 0 else 1)).value or "").strip() if title_col > 0 else ""
            region = str(ws.cell(r, (region_col if region_col > 0 else 1)).value or "").strip() if region_col > 0 else sheet_name.replace(" - PI ", "").split(" - ")[-1] if " - PI " in sheet_name else sheet_name

            # Skip if email already imported
            existing = db.query(Contact).filter(
                Contact.email == email, Contact.outreach_status != ""
            ).first()
            if existing:
                skipped += 1
                continue

            # Find or create Firm
            firm = None
            if firm_name:
                firm = db.query(Firm).filter(Firm.name == firm_name).first()
                if not firm:
                    firm = Firm(name=firm_name)
                    db.add(firm)
                    db.flush()

            # Find existing contact or create
            contact = db.query(Contact).filter(Contact.email == email).first()
            if contact:
                # Update with outreach info
                contact.outreach_status = "not_contacted"
                contact.outreach_region = region
                contact.outreach_notes = ""
                if title and not contact.title:
                    contact.title = title
                if phone and not contact.phone:
                    contact.phone = phone
            else:
                contact = Contact(
                    name=name,
                    email=email,
                    phone=phone[:50] if phone else "",
                    firm=firm_name,
                    firm_id=firm.id if firm else None,
                    title=title[:255] if title else "",
                    outreach_status="not_contacted",
                    outreach_region=region,
                    tags=["PI Law Firm", "Outreach"],
                )
                db.add(contact)
            created += 1

        db.commit()

    return {"ok": True, "created": created, "skipped": skipped, "total": created + skipped}


def _find_col(headers: list[str], *names: str) -> int:
    """Find the first column index matching any of the header names."""
    for n in names:
        for i, h in enumerate(headers):
            if n in h:
                return i + 1
    return -1


# ── Outreach Target List ──

@router.get("/targets")
def list_targets(
    region: str = Query(default=""),
    status: str = Query(default=""),
    search: str = Query(default=""),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    q = db.query(Contact).filter(Contact.outreach_status != "")
    if region:
        q = q.filter(Contact.outreach_region == region)
    if status:
        q = q.filter(Contact.outreach_status == status)
    if search:
        q = q.filter(
            Contact.name.ilike(f"%{search}%") |
            Contact.firm.ilike(f"%{search}%")
        )
    targets = q.order_by(Contact.outreach_region, Contact.name).all()
    return [
        {
            "id": t.id,
            "name": t.name,
            "email": t.email,
            "firm": t.firm,
            "title": t.title,
            "phone": t.phone,
            "region": t.outreach_region,
            "status": t.outreach_status,
            "notes": t.outreach_notes or "",
            "last_contacted_at": t.last_contacted_at.isoformat() if t.last_contacted_at else None,
        }
        for t in targets
    ]


@router.get("/regions")
def list_regions(
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    regions = db.query(Contact.outreach_region).filter(
        Contact.outreach_status != ""
    ).distinct().order_by(Contact.outreach_region).all()
    return [r[0] for r in regions if r[0]]


# ── Send Email ──

class SendOneRequest(BaseModel):
    contact_id: int


@router.post("/send-one")
async def send_one(
    data: SendOneRequest,
    _admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    if not is_smtp_configured():
        raise HTTPException(503, "SMTP not configured")

    contact = db.query(Contact).filter(Contact.id == data.contact_id).first()
    if not contact or not contact.email:
        raise HTTPException(404, "Contact not found or no email")

    # Personalize with AI
    body = PI_EMAIL_BODY.replace("{contact_name}", contact.name)
    body = body.replace("{firm}", contact.firm or "your firm")

    try:
        body = await personalize_with_ai(
            contact.name, contact.firm or "", contact.title or "", body
        )
    except Exception:
        pass  # use unpersonalized version

    subject = PI_EMAIL_SUBJECT.replace("{contact_name}", contact.name)

    success, error = await send_email(
        to_email=contact.email,
        to_name=contact.name,
        subject=subject,
        body=body,
    )

    if success:
        contact.outreach_status = "emailed"
        contact.last_contacted_at = datetime.now(timezone.utc)
        db.commit()
        return {"ok": True, "contact_id": contact.id, "status": "sent"}
    else:
        contact.outreach_status = "failed"
        contact.outreach_notes = f"Send error: {error[:200]}"
        db.commit()
        return {"ok": False, "contact_id": contact.id, "error": error}


@router.post("/send-batch")
async def send_batch(
    region: str = Query(default=""),
    status: str = Query(default="not_contacted"),
    _admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    """Send to all targets matching region + status filters."""
    if not is_smtp_configured():
        raise HTTPException(503, "SMTP not configured")

    q = db.query(Contact).filter(Contact.outreach_status != "", Contact.email != "")
    if region:
        q = q.filter(Contact.outreach_region == region)
    if status:
        q = q.filter(Contact.outreach_status == status)

    targets = q.limit(50).all()  # safety cap
    sent = 0
    failed = 0

    for contact in targets:
        body = PI_EMAIL_BODY.replace("{contact_name}", contact.name)
        body = body.replace("{firm}", contact.firm or "your firm")
        try:
            body = await personalize_with_ai(contact.name, contact.firm or "", contact.title or "", body)
        except Exception:
            pass

        subject = PI_EMAIL_SUBJECT.replace("{contact_name}", contact.name)
        success, error = await send_email(contact.email, contact.name, subject, body)

        if success:
            contact.outreach_status = "emailed"
            contact.last_contacted_at = datetime.now(timezone.utc)
            sent += 1
        else:
            contact.outreach_status = "failed"
            contact.outreach_notes = f"Send error: {error[:200]}"
            failed += 1
        db.commit()
        await asyncio.sleep(0.3)

    return {"ok": True, "sent": sent, "failed": failed}


@router.post("/mark/{contact_id}")
def mark_status(
    contact_id: int,
    status: str = Query(...),
    _admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    contact = db.query(Contact).filter(Contact.id == contact_id).first()
    if not contact:
        raise HTTPException(404, "Contact not found")
    contact.outreach_status = status
    db.commit()
    return {"ok": True, "contact_id": contact_id, "status": status}
